"""
MacDirScope

Description:
    A utility for macOS to scan a directory, extract rich filesystem and
    extended metadata, and export the results into a formatted Excel file.

    The program executes the following numbered steps:
      1. Checks if the system is macOS and has the 'mdls' command-line tool.
      2. Prompts the user to select an input directory to scan via a native dialog.
      3. Prompts the user for a save location for the output Excel file.
      4. Displays a configuration dialog allowing the user to optionally include
         media resolution and playback duration columns.
      5. Pre-scans the directory to count items and efficiently pre-computes all
         directory sizes for performance optimization.
      6. Displays a progress bar and begins processing every file and folder.
      7. For each item, extracts standard filesystem data (size, timestamps, status)
         and extended macOS metadata (Finder Tags, Kind, Resolution, Duration)
         using Spotlight metadata ('mdls').
      8. Writes collected data row-by-row into an Excel worksheet.
      9. Formats the Excel file with proportional column widths, date styling,
         numeric precision, and a frozen header row.
     10. Displays a final completion report summarizing scan statistics.

Usage:
    - Ensure required libraries are installed:
          pip install openpyxl
    - Run the script from a terminal:
          python mac_dir_scope.py

Author: Vitalii Starosta
GitHub: https://github.com/sztaroszta
License: MIT
"""

import os
import subprocess
import sys
from datetime import datetime
from tkinter import Tk, filedialog, messagebox, ttk, Label, Button, Toplevel, DoubleVar, BooleanVar
from typing import Tuple, Optional, List, Dict

from openpyxl import Workbook
from openpyxl.styles import NamedStyle, Font
from openpyxl.utils import get_column_letter

# --- Global Layout Configuration ---
HEADER_WIDTHS = {
    '#': 5,
    'Path': 25,
    'Size (KB)': 11,
    'Size (MB)': 11,
    'Creation Date': 19,
    'Last Modified': 19,
    'Is Hidden?': 10,
    'Tags': 10,
    'Kind': 15,
    'Resolution': 14,
    'Duration': 12,
    'File Type': 9,
}
LEVEL_COLUMN_WIDTH = 10

# Media extension filters for selective metadata extraction
IMAGE_EXTENSIONS = {
    '.jpg', '.jpeg', '.png', '.gif', '.bmp', '.tiff', '.tif',
    '.webp', '.heic', '.raw', '.cr2', '.nef'
}
VIDEO_EXTENSIONS = {
    '.mp4', '.mov', '.mkv', '.avi', '.m4v', '.wmv', '.flv', '.webm'
}
AUDIO_EXTENSIONS = {
    '.mp3', '.m4a', '.wav', '.flac', '.aac', '.aiff', '.ogg', '.wma'
}


# --- GUI Classes ---

class ScanOptionsWindow:
    """
    A GUI dialog prompting the user to select optional metadata columns
    (such as image/video resolution and media duration) prior to scanning.
    """

    def __init__(self):
        """Initializes and displays the scan configuration modal dialog."""
        self.root = Tk()
        self.root.title("Scan Configuration")
        self.root.geometry("460x220")
        self.root.resizable(False, False)

        self.resolution_var = BooleanVar(value=False)
        self.duration_var = BooleanVar(value=False)
        self.confirmed = False

        self.setup_widgets()

        self.root.protocol("WM_DELETE_WINDOW", self.on_cancel)
        self.root.lift()
        self.root.attributes('-topmost', True)
        self.root.mainloop()

    def setup_widgets(self):
        """Creates and arranges the configuration widgets within the window."""
        main_frame = ttk.Frame(self.root, padding="20 15 20 15")
        main_frame.pack(fill='both', expand=True)

        title_label = Label(
            main_frame,
            text="Additional Metadata Options",
            font=('Arial', 13, 'bold')
        )
        title_label.pack(anchor='w', pady=(0, 5))

        description_label = Label(
            main_frame,
            text="Select optional media columns to extract. Leave unchecked to skip.",
            font=('Arial', 10),
            fg='#555555'
        )
        description_label.pack(anchor='w', pady=(0, 15))

        options_frame = ttk.Frame(main_frame)
        options_frame.pack(fill='x', pady=5)

        res_check = ttk.Checkbutton(
            options_frame,
            text="Include Image & Video Resolution (e.g., 1920x1080)",
            variable=self.resolution_var
        )
        res_check.pack(anchor='w', pady=4)

        dur_check = ttk.Checkbutton(
            options_frame,
            text="Include Media Duration (Audio & Video)",
            variable=self.duration_var
        )
        dur_check.pack(anchor='w', pady=4)

        button_frame = ttk.Frame(main_frame)
        button_frame.pack(fill='x', pady=(15, 0))

        proceed_button = Button(
            button_frame,
            text="Proceed",
            command=self.on_proceed,
            width=12,
            default='active'
        )
        proceed_button.pack(side='right', padx=(5, 0))

        cancel_button = Button(
            button_frame,
            text="Cancel",
            command=self.on_cancel,
            width=12
        )
        cancel_button.pack(side='right')

    def on_proceed(self):
        """Stores the confirmed status and dismisses the dialog."""
        self.confirmed = True
        self.root.destroy()

    def on_cancel(self):
        """Dismisses the dialog without confirming."""
        self.confirmed = False
        self.root.destroy()


class ProgressWindow:
    """
    A GUI window to display the progress of a long-running task.
    Features a progress bar, status description, and processed item counter.
    """

    def __init__(self, total_items: int):
        """
        Initializes the ProgressWindow.

        Args:
            total_items (int): The total number of items to be processed.
        """
        self.root = Tk()
        self.root.withdraw()
        self.progress_toplevel = Toplevel(self.root)

        self.window = self.progress_toplevel
        self.window.title("Processing Directory...")
        self.window.geometry("600x150")
        self.window.resizable(False, False)
        self.window.transient()
        self.window.grab_set()

        self.total_items = total_items
        self.setup_widgets()

        self.window.lift()
        self.window.attributes('-topmost', True)

    def setup_widgets(self):
        """Creates and arranges all widgets within the progress window."""
        title_label = Label(self.window, text="Extracting Directory Metadata", font=('Arial', 13, 'bold'))
        title_label.pack(pady=10)

        self.status_label = Label(self.window, text="Initializing...")
        self.status_label.pack(pady=5)

        self.progress_var = DoubleVar()
        self.progress_bar = ttk.Progressbar(self.window, length=350, variable=self.progress_var, maximum=100)
        self.progress_bar.pack(pady=10)

        self.progress_label = Label(self.window, text="0 / 0 items processed")
        self.progress_label.pack(pady=5)

        self.cancel_button = Button(self.window, text="Run in Background", command=self.minimize_window)
        self.cancel_button.pack(pady=5)

    def update_progress(self, processed: int, status: str = ""):
        """
        Updates the progress bar percentage and status labels.

        Args:
            processed (int): The number of items processed so far.
            status (str, optional): A brief description of the current task.
        """
        progress_percent = (processed / self.total_items * 100) if self.total_items > 0 else 0
        self.progress_var.set(progress_percent)
        if status:
            self.status_label.config(text=status)
        self.progress_label.config(text=f"{processed} / {self.total_items} items processed")
        self.window.update()

    def minimize_window(self):
        """Minimizes the progress window to the system dock."""
        self.window.iconify()

    def close(self):
        """Destroys the progress window and terminates its Tkinter root."""
        try:
            self.root.destroy()
        except:
            pass


class CompletionReportWindow:
    """
    A GUI dialog that presents a statistical summary of the scan results.
    Runs its own mainloop to act as a blocking dialog upon process completion.
    """

    def __init__(self, stats: dict):
        """
        Initializes and displays the completion report dialog.

        Args:
            stats (dict): Dictionary containing summary metrics of the scan.
        """
        self.window = Tk()
        self.window.title("Processing Complete")
        self.window.geometry("550x350")
        self.window.resizable(False, False)

        self.stats = stats
        self.setup_widgets()

        self.window.lift()
        self.window.attributes('-topmost', True)
        self.window.mainloop()

    def setup_widgets(self):
        """Creates and arranges summary elements within the report window."""
        title_label = Label(self.window, text="✓ Processing Complete", font=('Arial', 15, 'bold'), fg='green')
        title_label.pack(pady=15)

        stats_frame = ttk.Frame(self.window)
        stats_frame.pack(pady=10, padx=20, fill='both', expand=True)

        stats_text = f"""Directory Metadata Extraction Results:

Directory Scanned: {self.stats.get('directory', 'N/A')}
Items Processed: {self.stats.get('processed_items', 0):,}
   • Directories: {self.stats.get('directories', 0):,}
   • Files: {self.stats.get('files', 0):,}
Max Depth: {self.stats.get('max_levels', 0)} levels
Total Size: {self.stats.get('total_size_mb', 0):.2f} MB
Output File: {self.stats.get('output_file', 'N/A')}
Processing Time: {self.stats.get('duration', 'N/A')}"""

        stats_label = Label(stats_frame, text=stats_text, justify='left', font=('Courier', 12))
        stats_label.pack(pady=10)

        button_frame = ttk.Frame(self.window)
        button_frame.pack(pady=15)

        close_button = Button(button_frame, text="Close", command=self.window.destroy, width=15)
        close_button.pack(side='right', padx=5)

        if self.stats.get('output_file'):
            open_button = Button(button_frame, text="Open File Location", command=self.open_file_location, width=15)
            open_button.pack(side='left', padx=5)

    def open_file_location(self):
        """Reveals the output Excel file in the macOS Finder."""
        try:
            output_file = self.stats.get('output_file')
            if output_file and os.path.exists(output_file):
                subprocess.run(['open', '-R', output_file])
        except Exception as e:
            print(f"Could not open file location: {e}")


# --- Metadata Extraction Functions ---

def check_mdls_availability() -> bool:
    """
    Checks if the macOS 'mdls' command-line tool is available and executable.

    Returns:
        bool: True if mdls is present and responsive, False otherwise.
    """
    try:
        subprocess.run(['mdls', '--help'], capture_output=True, check=True)
        return True
    except (subprocess.CalledProcessError, FileNotFoundError):
        return False


def get_file_tags(path: str) -> str:
    """
    Retrieves Finder user tags for a specified path using 'mdls'.

    Args:
        path (str): Full path to the file or directory.

    Returns:
        str: Comma-separated list of tags, or an empty string if none exist.
    """
    try:
        result = subprocess.run(
            ['mdls', '-name', 'kMDItemUserTags', '-raw', path],
            capture_output=True, text=True, timeout=10
        )
        return process_tags(result.stdout.strip())
    except:
        return ""


def process_tags(tags_str: str) -> str:
    """
    Parses and cleans the raw tag string returned by 'mdls'.

    Args:
        tags_str (str): The raw stdout response from the mdls query.

    Returns:
        str: Cleaned, comma-separated tag string.
    """
    if not tags_str or tags_str == "(null)":
        return ""
    tags_str = tags_str.strip('()')
    return ', '.join([tag.strip().strip('"') for tag in tags_str.split(',') if tag.strip()]) if tags_str else ""


def get_file_kind(path: str) -> str:
    """
    Retrieves the macOS 'Kind' descriptor for a file (e.g., 'JPEG image', 'Folder').

    Args:
        path (str): Full path to the file or directory.

    Returns:
        str: Descriptive Kind label, or an empty string if unavailable.
    """
    try:
        result = subprocess.run(
            ['mdls', '-name', 'kMDItemKind', '-raw', path],
            capture_output=True, text=True, timeout=10
        )
        kind = result.stdout.strip()
        return kind if kind != "(null)" else ""
    except:
        return ""


def get_media_resolution(path: str) -> str:
    """
    Retrieves pixel dimensions (Width x Height) for image or video files using 'mdls'.

    Args:
        path (str): Full path to the media file.

    Returns:
        str: Formatted resolution string (e.g., '1920x1080'), or an empty string.
    """
    _, ext = os.path.splitext(path.lower())
    if ext not in IMAGE_EXTENSIONS and ext not in VIDEO_EXTENSIONS:
        return ""

    try:
        result = subprocess.run(
            ['mdls', '-name', 'kMDItemPixelWidth', '-name', 'kMDItemPixelHeight', path],
            capture_output=True, text=True, timeout=5
        )
        width, height = None, None
        for line in result.stdout.splitlines():
            if 'kMDItemPixelWidth' in line and '=' in line:
                val = line.split('=')[-1].strip()
                if val and val != "(null)":
                    width = val
            elif 'kMDItemPixelHeight' in line and '=' in line:
                val = line.split('=')[-1].strip()
                if val and val != "(null)":
                    height = val

        if width and height:
            return f"{width}x{height}"
        return ""
    except:
        return ""


def get_media_duration(path: str) -> str:
    """
    Retrieves playback duration formatted as 'HH:MM:SS' or 'MM:SS' for media files.

    Args:
        path (str): Full path to the audio or video file.

    Returns:
        str: Formatted duration string, or an empty string if unavailable.
    """
    _, ext = os.path.splitext(path.lower())
    if ext not in VIDEO_EXTENSIONS and ext not in AUDIO_EXTENSIONS:
        return ""

    try:
        result = subprocess.run(
            ['mdls', '-name', 'kMDItemDurationSeconds', '-raw', path],
            capture_output=True, text=True, timeout=5
        )
        duration_raw = result.stdout.strip()
        if not duration_raw or duration_raw == "(null)":
            return ""

        seconds = float(duration_raw)
        total_seconds = int(round(seconds))
        hours = total_seconds // 3600
        minutes = (total_seconds % 3600) // 60
        secs = total_seconds % 60

        if hours > 0:
            return f"{hours:02d}:{minutes:02d}:{secs:02d}"
        return f"{minutes:02d}:{secs:02d}"
    except:
        return ""


# --- Filesystem and Excel Processing Functions ---

def precompute_directory_sizes(root_path: str) -> Dict[str, int]:
    """
    Performs an initial walk of the directory tree to calculate the total size
    of every subdirectory in bytes, aggregating sizes upwards.

    Args:
        root_path (str): Top-level directory path to begin the scan from.

    Returns:
        Dict[str, int]: Mapping of directory paths to their aggregated sizes in bytes.
    """
    dir_sizes = {}
    for root, _, files in os.walk(root_path):
        try:
            size = sum(
                os.path.getsize(os.path.join(root, f))
                for f in files
                if not os.path.islink(os.path.join(root, f))
            )
            dir_sizes[root] = size
        except OSError:
            dir_sizes[root] = 0

    # Aggregate child directory sizes into parent directories
    for path in sorted(dir_sizes.keys(), key=len, reverse=True):
        parent = os.path.dirname(path)
        if parent != path and parent in dir_sizes:
            dir_sizes[parent] += dir_sizes[path]

    return dir_sizes


def get_file_info(
    path: str,
    directory_sizes: Dict[str, int],
    include_resolution: bool = False,
    include_duration: bool = False
) -> Optional[Tuple]:
    """
    Gathers standard filesystem attributes and optional extended metadata for an item.

    Args:
        path (str): Full path to the file or directory.
        directory_sizes (Dict[str, int]): Pre-computed directory size lookup table.
        include_resolution (bool): Whether to query resolution for media files.
        include_duration (bool): Whether to query duration for media files.

    Returns:
        Optional[Tuple]: Tuple containing item metadata fields and an optional
                         fields list, or None if the item is inaccessible.
    """
    try:
        stat_info = os.stat(path)
        created = datetime.fromtimestamp(stat_info.st_birthtime)
        modified = datetime.fromtimestamp(stat_info.st_mtime)

        if os.path.isdir(path):
            size_in_bytes = directory_sizes.get(path, 0)
            size_kb = size_in_bytes / 1024
            size_mb = size_in_bytes / (1024 * 1024)
            file_type = "Folder"
            resolution = "" if include_resolution else None
            duration = "" if include_duration else None
        else:
            size_in_bytes = stat_info.st_size
            size_kb = size_in_bytes / 1024
            size_mb = size_in_bytes / (1024 * 1024)
            _, ext = os.path.splitext(os.path.basename(path))
            file_type = ext[1:] if ext else "File"
            resolution = get_media_resolution(path) if include_resolution else None
            duration = get_media_duration(path) if include_duration else None

        basename = os.path.basename(path)
        hidden = "hidden" if basename.startswith('.') else "temporary" if basename.startswith('~$') else "visible"
        tags = get_file_tags(path)
        kind = get_file_kind(path)

        optional_fields = []
        if include_resolution:
            optional_fields.append(resolution)
        if include_duration:
            optional_fields.append(duration)

        return (created, modified, size_kb, size_mb, file_type, hidden, tags, kind, optional_fields)
    except:
        return None


def get_path_levels(path: str) -> List[str]:
    """
    Deconstructs a filesystem path into its constituent directory levels.

    Args:
        path (str): The full filesystem path.

    Returns:
        List[str]: Directory and file names split by the system path separator.
    """
    return [level for level in path.split(os.sep) if level]


def count_files_and_max_levels(starting_directory: str) -> Tuple[int, int]:
    """
    Executes a fast pre-scan to compute total item count and maximum directory depth.

    Args:
        starting_directory (str): Root path to evaluate.

    Returns:
        Tuple[int, int]: Total item count and maximum directory level depth.
    """
    total_items, max_levels = 0, 0
    try:
        for root, dirs, files in os.walk(starting_directory):
            for item in dirs + files:
                total_items += 1
                path = os.path.join(root, item)
                max_levels = max(max_levels, len(get_path_levels(path)))
    except:
        pass
    return total_items, max_levels


def setup_worksheet_headers(
    worksheet,
    max_levels: int,
    include_resolution: bool = False,
    include_duration: bool = False
) -> List[str]:
    """
    Constructs and appends the header row into the active Excel worksheet.

    Args:
        worksheet: The target openpyxl worksheet instance.
        max_levels (int): Maximum depth used to generate 'Level N' column titles.
        include_resolution (bool): Whether the Resolution header is included.
        include_duration (bool): Whether the Duration header is included.

    Returns:
        List[str]: Complete list of ordered column header titles.
    """
    headers = ['#', 'Path', 'Size (KB)', 'Size (MB)', 'Creation Date', 'Last Modified', 'Is Hidden?', 'Tags', 'Kind']
    if include_resolution:
        headers.append('Resolution')
    if include_duration:
        headers.append('Duration')
    headers.append('File Type')
    headers.extend([f'Level {i+1}' for i in range(max_levels)])
    worksheet.append(headers)
    return headers


def format_worksheet(worksheet, headers: List[str]):
    """
    Applies custom column widths, number and date formatting, and freezes headers.

    Args:
        worksheet: The openpyxl worksheet instance to format.
        headers (List[str]): List of column header names used for dimension lookup.
    """
    date_style = NamedStyle(name='datetime', number_format='YYYY-MM-DD HH:MM:SS')

    # Apply column widths based on header mappings
    for idx, header in enumerate(headers, start=1):
        col_letter = get_column_letter(idx)
        if header in HEADER_WIDTHS:
            worksheet.column_dimensions[col_letter].width = HEADER_WIDTHS[header]
        elif 'Level' in header:
            worksheet.column_dimensions[col_letter].width = LEVEL_COLUMN_WIDTH

    # Header row formatting
    for cell in worksheet[1]:
        cell.font = Font(bold=True)

    worksheet.freeze_panes = 'C2'
    worksheet.auto_filter.ref = worksheet.dimensions

    # Apply datetime formatting to Creation Date (col 5) and Last Modified (col 6)
    for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row, min_col=5, max_col=6):
        for cell in row:
            cell.style = date_style

    # Apply two-decimal precision to Size (KB) and Size (MB) (cols 3 & 4)
    for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row, min_col=3, max_col=4):
        for cell in row:
            cell.number_format = '0.00'


def generate_excel(
    starting_directory: str,
    save_path: str,
    include_resolution: bool = False,
    include_duration: bool = False
) -> Tuple[bool, dict]:
    """
    Coordinates directory scanning, metadata retrieval, and Excel file generation.

    Args:
        starting_directory (str): Root directory path to scan.
        save_path (str): Destination file path for the exported .xlsx workbook.
        include_resolution (bool): Whether to extract media resolution.
        include_duration (bool): Whether to extract media playback duration.

    Returns:
        Tuple[bool, dict]: Success flag and dictionary of operational metrics.
    """
    start_time = datetime.now()
    total_items, max_levels = count_files_and_max_levels(starting_directory)
    if total_items == 0:
        return False, {}

    print("Pre-computing directory sizes for performance...")
    directory_sizes = precompute_directory_sizes(starting_directory)
    print("Pre-computation complete. Starting main processing...")

    progress_window = ProgressWindow(total_items)
    progress_window.update_progress(0, "Setting up...")

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = 'Directory Info'
    headers = setup_worksheet_headers(worksheet, max_levels, include_resolution, include_duration)

    stats = {
        'directory': starting_directory, 'output_file': save_path,
        'total_items': total_items, 'max_levels': max_levels,
        'processed_items': 0, 'directories': 0, 'files': 0, 'errors': 0,
        'total_size_mb': 0, 'duration': '0s'
    }

    row_number, processed_items = 1, 0
    try:
        for root, dirs, files in os.walk(starting_directory):
            all_items = [(d, True) for d in dirs] + [(f, False) for f in files]
            for item_name, is_dir in all_items:
                current_path = os.path.join(root, item_name)
                file_info = get_file_info(current_path, directory_sizes, include_resolution, include_duration)

                if file_info:
                    created, mod, size_kb, size_mb, ftype, hidden, tags, kind, optional_fields = file_info
                    path_levels = get_path_levels(current_path)
                    row_data = [
                        row_number, current_path, size_kb, size_mb, created, mod,
                        hidden, tags, kind, *optional_fields, ftype, *path_levels
                    ]
                    worksheet.append(row_data)
                    row_number += 1
                    stats['directories' if is_dir else 'files'] += 1
                else:
                    stats['errors'] += 1

                processed_items += 1
                stats['processed_items'] = processed_items
                if processed_items % 10 == 0:
                    progress_window.update_progress(processed_items, f"Processing: {os.path.basename(current_path)}")

    except Exception as e:
        print(f"An error occurred during file processing: {e}")
        stats['errors'] += 1
        progress_window.close()
        return False, stats

    # Retrieve total root directory size from pre-computed metrics
    stats['total_size_mb'] = directory_sizes.get(starting_directory, 0) / (1024 * 1024)

    progress_window.update_progress(processed_items, "Formatting and saving...")
    format_worksheet(worksheet, headers)

    try:
        workbook.save(save_path)
        stats['duration'] = str(datetime.now() - start_time).split('.')[0]
        progress_window.close()
        return True, stats
    except Exception as e:
        print(f"Error saving Excel file: {e}")
        progress_window.close()
        return False, stats


# --- Main Application Logic ---

def prompt_scan_options() -> Optional[Tuple[bool, bool]]:
    """
    Displays the scan options dialog and retrieves user preferences.

    Returns:
        Optional[Tuple[bool, bool]]: A tuple of (include_resolution, include_duration),
                                     or None if the user cancelled the dialog.
    """
    dialog = ScanOptionsWindow()
    if not dialog.confirmed:
        return None
    return dialog.resolution_var.get(), dialog.duration_var.get()


def get_directory_and_save_path() -> Tuple[Optional[str], Optional[str]]:
    """
    Prompts the user for directory and output paths using native dialogs.

    Returns:
        Tuple[Optional[str], Optional[str]]: (starting_directory, save_path),
                                             or (None, None) if cancelled.
    """
    root_dir = Tk()
    root_dir.withdraw()
    starting_directory = filedialog.askdirectory(title="Select the directory to scan")
    root_dir.destroy()
    if not starting_directory:
        return None, None

    root_save = Tk()
    root_save.withdraw()
    directory_name = os.path.basename(starting_directory)
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    default_filename = f"{directory_name}_{timestamp}.xlsx"
    save_path = filedialog.asksaveasfilename(
        title="Save Excel file as...", initialfile=default_filename,
        defaultextension=".xlsx", filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")]
    )
    root_save.destroy()
    if not save_path:
        return None, None

    return starting_directory, save_path


def main():
    """Main execution entry point for MacDirScope."""
    print("macOS Directory Metadata Extractor")
    print("=" * 40)

    if not check_mdls_availability():
        root_err = Tk()
        root_err.withdraw()
        messagebox.showerror("Dependency Error", "This script requires macOS and the 'mdls' command.")
        root_err.destroy()
        sys.exit(1)

    starting_directory, save_path = get_directory_and_save_path()

    if not starting_directory or not save_path:
        print("Operation cancelled by user.")
        sys.exit(0)

    # Prompt user for optional column inclusion
    options = prompt_scan_options()
    if options is None:
        print("Operation cancelled by user.")
        sys.exit(0)

    include_resolution, include_duration = options

    print(f"Scanning directory: {starting_directory}")
    print(f"Output file: {save_path}")
    print(f"Options: Resolution={include_resolution}, Duration={include_duration}")

    success, stats = generate_excel(
        starting_directory,
        save_path,
        include_resolution=include_resolution,
        include_duration=include_duration
    )

    if success:
        print("\nOperation completed successfully!")
        CompletionReportWindow(stats)
    else:
        print("\nOperation failed. Please check the error messages above.")
        root_err = Tk()
        root_err.withdraw()
        messagebox.showerror("Error", "Processing failed. Please check the console for details.")
        root_err.destroy()
        sys.exit(1)


if __name__ == "__main__":
    main()